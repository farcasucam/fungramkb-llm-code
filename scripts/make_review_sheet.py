"""Expert-review sheet for the concepts authored for this study (validation level 1 of the project script).

One row per authored concept (provenance ``authored-*``): taxonomy, lemmas, gloss, the COREL meaning postulate and
its verbalisation (what the models actually read in the G2 context), and how many items of the main study's
analysis set depend on the concept. Columns for the reviewer: taxonomy / postulate / lemmas OK (1/0), corrected
value, comment. Rows are ordered by the number of dependent items, so the most influential concepts come first.

    python scripts/make_review_sheet.py --kb data/processed/fungramkb.json \\
        --bench data/processed/fgkb_reason_main.jsonl --out docs/review_authored_concepts.xlsx
    python scripts/make_review_sheet.py --summarise docs/review_authored_concepts.xlsx   # after the review
"""

from __future__ import annotations

import argparse
import collections
import re

from fgkb_llm.bench.contrast import contrast_subset
from fgkb_llm.bench.schema import read_jsonl
from fgkb_llm.corel.verbalize import verbalise_fact, verbalise_isa
from fgkb_llm.kb.loaders import load_json
from fgkb_llm.reasoner.asp import Reasoner

REVIEW_COLS = ["taxonomy_ok (1/0)", "postulate_ok (1/0)", "lemmas_ok (1/0)", "corrected postulate / parent / lemma",
               "comment"]
INSTRUCTIONS = [
    ("Purpose", ("Expert validation of the concepts authored for the study (they are about a quarter of the KB). "
                "The paper reports the share approved and, if needed, results restricted to items that use only "
                "concepts of the original export.")),
    ("taxonomy_ok", ("1 if the parent (IS-A) is correct in FunGramKB terms; 0 otherwise (write the right parent in the "
                    "'corrected' column).")),
    ("postulate_ok", ("1 if the meaning postulate is correct COREL and states correct knowledge (strict + vs "
                     "defeasible *, roles, fillers); 0 otherwise. 'verbalised (EN)' is what the models read.")),
    ("lemmas_ok", "1 if the English and Spanish lemmas are right for the concept."),
    ("items", ("Number of test items of the main study (contrast subset, EN+ES) that use the concept: as the "
              "queried concept, as an ancestor in the inference chain or as a filler. Rows with many items matter "
              "most; rows with 0 items can be reviewed quickly.")),
    ("Second rater", ("For inter-rater agreement, a second expert fills the same three 1/0 columns for the first 40 "
                     "rows in a copy of this file (review_authored_concepts_rater2.xlsx).")),
]


def usage_counts(bench: str) -> collections.Counter:
    pool = [i for i in read_jsonl(bench) if i.task == "deep" and i.split in ("test_seen", "test_unseen", "dev")]
    items = [i for i in contrast_subset(pool) if i.split in ("test_seen", "test_unseen")]
    n = collections.Counter()
    for it in items:
        used = set(it.focus_concepts)
        for f in it.meta.get("required_facts") or []:
            used |= set(re.findall(r"[+$#][A-Z0-9_]+", f))
        prop = it.meta.get("prop") or []
        used |= {p for p in prop if isinstance(p, str) and p[:1] in "+$"}
        n.update(used)
    return n


def build(kb_path: str, bench: str, out: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    kb = load_json(kb_path)
    r = Reasoner(kb)
    n = usage_counts(bench)
    rows = []
    for cid, con in kb.concepts.items():
        if not str((con.meta or {}).get("provenance", "")).startswith("authored"):
            continue
        facts = [verbalise_fact(kb, f, "en") for f in r.facts_by_concept.get(cid, [])]
        isa = [verbalise_isa(kb, cid, p, "en") for p in sorted(r.parents.get(cid, ())) if not p.startswith("#")]
        rows.append({"concept": cid, "type": con.semantic_type, "items": n.get(cid, 0),
                     "parents": ", ".join(con.parents), "lemma_en": ", ".join(kb.lemmas(cid, "en")),
                     "lemma_es": ", ".join(kb.lemmas(cid, "es")), "gloss": (con.glosses or {}).get("en", ""),
                     "meaning_postulate": con.meaning_postulate or "",
                     "verbalised (EN)": "\n".join(isa + facts)})
    rows.sort(key=lambda x: (-x["items"], x["type"], x["concept"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "concepts"
    head = list(rows[0]) + REVIEW_COLS
    ws.append(head)
    for row in rows:
        ws.append([row[k] for k in rows[0]] + [""] * len(REVIEW_COLS))
    widths = {"concept": 24, "type": 8, "items": 7, "parents": 24, "lemma_en": 16, "lemma_es": 16, "gloss": 30,
              "meaning_postulate": 60, "verbalised (EN)": 60}
    for k, col in enumerate(head, start=1):
        letter = ws.cell(row=1, column=k).column_letter
        ws.column_dimensions[letter].width = widths.get(col, 16 if "ok" in col else 40)
        ws.cell(row=1, column=k).font = Font(bold=True)
        if col in REVIEW_COLS:
            ws.cell(row=1, column=k).fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("instructions", 0)
    for k, v in INSTRUCTIONS:
        info.append([k, v])
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out)
    by_type = collections.Counter(x["type"] for x in rows)
    print(f"{len(rows)} authored concepts -> {out} ({dict(by_type)}); used by test items: "
          f"{sum(1 for x in rows if x['items'])}")


def summarise(path: str):
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True)["concepts"]
    rows = list(ws.iter_rows(values_only=True))
    head = list(rows[0])
    data = [dict(zip(head, r)) for r in rows[1:]]
    for col in REVIEW_COLS[:3]:
        vals = [str(d[col]).strip() for d in data if d[col] not in (None, "")]
        ok = sum(v == "1" for v in vals)
        print(f"{col}: {ok}/{len(vals)} approved ({100 * ok / max(1, len(vals)):.1f} %), {len(data) - len(vals)} blank")
    weighted = sum(d["items"] or 0 for d in data if str(d[REVIEW_COLS[1]]).strip() == "0")
    print(f"test items using a concept whose postulate was rejected: {weighted} (concept-item uses)")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--kb", default="data/processed/fungramkb.json")
    ap.add_argument("--bench", default="data/processed/fgkb_reason_main.jsonl")
    ap.add_argument("--out", default="docs/review_authored_concepts.xlsx")
    ap.add_argument("--summarise", default=None)
    a = ap.parse_args(argv)
    if a.summarise:
        summarise(a.summarise)
    else:
        build(a.kb, a.bench, a.out)


if __name__ == "__main__":
    main()
