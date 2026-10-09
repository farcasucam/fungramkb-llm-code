"""Expert-rating sheet for the generated questions (NLG audit, validation level 2 of the project script).

Turns docs/nlg_audit_sample_main.tsv (120 items stratified by block) into an Excel sheet that shows, next to each
EN/ES question, what the question is meant to express according to the knowledge base: the queried predication in
COREL-like notation, the reasoning chain behind the gold answer (IS-A steps and the supporting postulate, quoted
from the meaning postulate), and for non-deep tasks the premise, hypothesis or options. The rater fills fluency
(1-5, EN and ES) and faithfulness (1/0).

    python scripts/make_nlg_audit_sheet.py --sample docs/nlg_audit_sample_main.tsv \\
        --bench data/processed/fgkb_reason_main.jsonl --kb data/processed/fungramkb.json \\
        --out docs/nlg_audit_sample_main.xlsx
    python scripts/make_nlg_audit_sheet.py --summarise docs/nlg_audit_sample_main.xlsx \\
        [--second docs/nlg_audit_sample_main_rater2.xlsx]
"""

from __future__ import annotations

import argparse
import ast
import csv
import json
import re
import statistics

from fgkb_llm.kb.loaders import load_json

RATE_COLS = ["fluency_en (1-5)", "fluency_es (1-5)", "faithful (1/0)", "comment"]
INSTRUCTIONS = [
    ("Purpose", ("Check that the questions generated from the knowledge base read naturally and say what the "
                "knowledge base says. The paper reports mean fluency and the share of faithful items.")),
    ("fluency_en / fluency_es", ("5 = natural, as a native speaker would write it; 4 = correct but slightly stiff; "
                                "3 = understandable with awkward wording (articles, word order, register); "
                                "2 = hard to understand or ungrammatical; 1 = unintelligible. Judge only the language, "
                                "not whether the content is true.")),
    ("faithful", ("1 if the question (in both languages) asks exactly the predication in 'queried predication' about "
                 "the concept in 'about' (same event, same participants and roles, same polarity), and any "
                 "background sentence states what 'scenario / premise' says; 0 if it asks something different "
                 "(wrong role, wrong participant, wrong polarity, ES and EN mean different things). Say what is "
                 "wrong in 'comment'.")),
    ("gold / chain", ("For information: the reasoner's answer and the inheritance chain behind it. You do not need to "
                     "check the gold; if you think it is wrong, write it in 'comment'.")),
    ("Which rows", ("The first 70 rows (deep blocks: deep_real, novel, exceptions, counterfactual, ablation, "
                   "undetermined, undet_control) are the ones the paper's main study uses: rate these. The last 50 "
                   "rows (entailment, consistency, multihop, grounding, procedural) belong to tasks not analysed in "
                   "this paper and are optional.")),
    ("Second rater", ("For agreement, a second expert rates the first 40 rows (deep blocks) in a copy named "
                     "nlg_audit_sample_main_rater2.xlsx.")),
]


def _predication(mp: str, evar: str) -> str:
    """The predication '(eN: ...)' with its prefix (+ or *) quoted from a meaning postulate."""
    key = f"({evar}:"
    i = mp.find(key)
    if i < 0:
        return ""
    depth, j = 0, i
    while j < len(mp):
        depth += mp[j] == "("
        depth -= mp[j] == ")"
        j += 1
        if depth == 0:
            break
    sign = mp[i - 1] if i > 0 and mp[i - 1] in "+*" else ""
    return sign + mp[i:j]


def describe(it: dict, kb) -> dict:
    meta = it.get("meta") or {}
    about = ", ".join(it.get("focus_concepts") or [])
    prop = meta.get("prop") or []
    queried = ""
    hops = [ast.literal_eval(t) for t in re.findall(r"\('[^)]*\)", it.get("pair_id") or "")]
    if len(hops) >= 2:  # composed two-step item: "X EV1 something that EV2 F2"
        (e1, r1a, r1b, f1), (e2, r2a, r2b, f2) = hops[:2]
        focus = (it.get("focus_concepts") or ["?"])[0]
        queried = (f"step 1: ({focus} as {r1a}) {e1} ({f1} as {r1b})\n"
                   f"step 2: ({f1} as {r2a}) {e2}" + (f" ({f2} as {r2b})" if f2 else "") +
                   "\n(the question composes both: '... something that ...')")
    elif len(prop) == 4:
        ev, srole, orole, filler = prop
        queried = f"({it.get('focus_concepts', ['?'])[0]} as {srole}) {ev}" + (f" ({filler} as {orole})" if filler else "")
    chain = []
    for step in it.get("trace") or []:
        if step and step[0] == "isa":
            chain.append(f"{step[1]} IS-A {step[2]}")
        elif step and step[0] == "postulate":
            cid, src = step[1], step[2]
            evar = src.split("/")[-1]
            patched = ((meta.get("kb_patch") or {}).get("concepts") or {}).get(cid) or {}
            con = kb.concepts.get(cid)
            mp = patched.get("meaning_postulate") or (con.meaning_postulate if con else "") or ""
            pred = _predication(mp, evar)
            chain.append(f"{cid} postulate {evar} ({step[3] if len(step) > 3 else ''}): {pred}")
    scenario = []
    if it.get("premise"):
        scenario.append("Premise: " + it["premise"])
    if it.get("hypothesis"):
        scenario.append("Hypothesis: " + it["hypothesis"])
    if it.get("options"):
        scenario.append("Options: " + " | ".join(it["options"]))
    if meta.get("removed"):
        scenario.append(f"Removed from the KB for this item: {meta['removed']}")
    patch = (meta.get("kb_patch") or {}).get("concepts") or {}
    for cid, spec in list(patch.items())[:4]:
        scenario.append(f"Patched {cid}: parents={spec.get('parents')} postulate={spec.get('meaning_postulate')}")
    return {"about": about, "queried predication": queried, "chain behind the gold": "\n".join(chain),
            "scenario / premise": "\n".join(scenario)}


def build(sample: str, bench: str, kb_path: str, out: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    kb = load_json(kb_path)
    with open(sample, encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh, delimiter="\t"))
    deep_blocks = ("deep_real", "novel", "exceptions", "counterfactual", "ablation", "undetermined", "undet_control")
    # the main study analyses the deep task only: its blocks first, the other tasks (optional) after them
    rows.sort(key=lambda r: (r["block"] not in deep_blocks, r["block"]))
    ids = {r["id"] for r in rows}
    items = {}
    with open(bench, encoding="utf-8") as fh:
        for line in fh:
            it = json.loads(line)
            if it["id"] in ids:
                items[it["id"]] = it
    wb = Workbook()
    ws = wb.active
    ws.title = "questions"
    head = ["id", "block", "question_en", "question_es", "about", "queried predication", "scenario / premise",
            "gold", "chain behind the gold"] + RATE_COLS
    ws.append(head)
    for r in rows:
        d = describe(items[r["id"]], kb) if r["id"] in items else {}
        ws.append([r["id"], r["block"], r["question_en"], r["question_es"], d.get("about", ""),
                   d.get("queried predication", ""), d.get("scenario / premise", ""), r["gold"],
                   d.get("chain behind the gold", "")] + [""] * len(RATE_COLS))
    widths = {"id": 24, "block": 13, "question_en": 50, "question_es": 50, "about": 20, "queried predication": 36,
              "scenario / premise": 40, "gold": 12, "chain behind the gold": 55, "comment": 40}
    for k, col in enumerate(head, start=1):
        cell = ws.cell(row=1, column=k)
        ws.column_dimensions[cell.column_letter].width = widths.get(col, 14)
        cell.font = Font(bold=True)
        if col in RATE_COLS:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("instructions", 0)
    for k, v in INSTRUCTIONS:
        info.append([k, v])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out)
    print(f"{len(rows)} items -> {out}; with KB details: {len(items)}")


def _ratings(path: str) -> dict:
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True)["questions"]
    rows = list(ws.iter_rows(values_only=True))
    head = list(rows[0])
    return {r[0]: dict(zip(head, r)) for r in rows[1:] if r[0]}


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def summarise(path: str, second: str | None):
    a = _ratings(path)
    for col in RATE_COLS[:2]:
        vals = [x for x in (_num(r[col]) for r in a.values()) if x is not None]
        if vals:
            print(f"{col}: mean {statistics.mean(vals):.2f}, share >= 4: {sum(v >= 4 for v in vals) / len(vals):.2f} "
                  f"(n={len(vals)})")
    f = [x for x in (_num(r[RATE_COLS[2]]) for r in a.values()) if x is not None]
    if f:
        print(f"faithful: {sum(f):.0f}/{len(f)} ({100 * sum(f) / len(f):.1f} %)")
    if second:
        from sklearn.metrics import cohen_kappa_score

        b = _ratings(second)
        for col in RATE_COLS[:3]:
            pairs = [(_num(a[i][col]), _num(b[i][col])) for i in a if i in b]
            pairs = [(x, y) for x, y in pairs if x is not None and y is not None]
            if len(pairs) >= 5:
                x, y = zip(*pairs)
                w = None if col.startswith("faithful") else "quadratic"
                print(f"{col}: Cohen's kappa{' (quadratic)' if w else ''} = {cohen_kappa_score(x, y, weights=w):.2f} "
                      f"(n={len(pairs)})")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--sample", default="docs/nlg_audit_sample_main.tsv")
    ap.add_argument("--bench", default="data/processed/fgkb_reason_main.jsonl")
    ap.add_argument("--kb", default="data/processed/fungramkb.json")
    ap.add_argument("--out", default="docs/nlg_audit_sample_main.xlsx")
    ap.add_argument("--summarise", default=None)
    ap.add_argument("--second", default=None, help="second rater's copy, for Cohen's kappa")
    ap.add_argument("--v2-from", nargs=2, metavar=("FIRST_SHEET", "OUT"), default=None,
                    help="build the faithfulness re-rating sheet (two criteria, deep rows) from a rater's first sheet")
    ap.add_argument("--v2-summarise", nargs=2, metavar=("RATER1", "RATER2"), default=None)
    a = ap.parse_args(argv)
    if a.v2_from:
        build_v2(*a.v2_from)
    elif a.v2_summarise:
        summarise_v2(*a.v2_summarise)
    elif a.summarise:
        summarise(a.summarise, a.second)
    else:
        build(a.sample, a.bench, a.kb, a.out)



# ---- faithfulness re-rating (v2): two criteria, deep blocks only -----------------------------------------------
V2_COLS = ["faithful_structure (1/0)", "lexical_choice_ok (1/0)", "comment v2"]
V2_INSTRUCTIONS = [
    ("What to do", ("Rate ONLY the two yellow columns, for every row (70 items of the deep blocks analysed in the paper). "
                   "Fluency is already filled in from your first rating and does not need to be redone.")),
    ("faithful_structure", ("1 if the question has the STRUCTURE of 'queried predication': same event, same participants "
                           "in the same roles, same polarity, and the background sentence (if any) says what "
                           "'scenario / premise' says. Judge the concepts as given (do not penalise the word chosen for a "
                           "concept here). 0 if the event, a participant, a role or the polarity is wrong, or EN and ES "
                           "differ in structure.")),
    ("lexical_choice_ok", ("1 if the English and Spanish words chosen for every concept identify that concept "
                          "adequately; 0 if a word names a different or narrower/wider concept (e.g. 'confinement' for "
                          "+IMPRISONMENT_00, 'refugee' for +REFUGE_00). Grammar and fluency do not count here.")),
    ("Composed items", ("Rows shaded salmon are two-step questions ('... something that ...'): 'queried predication' shows "
                       "both steps and the question is faithful if it expresses both.")),
    ("Independence", "Each rater fills his own file without looking at the other's."),
]


def build_v2(old_path: str, out: str):
    """Faithfulness re-rating sheet from a rater's first sheet: deep rows only, fluency kept, two new criteria."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    src = load_workbook(old_path)["questions"]
    head = [c.value for c in src[1]]
    deep = ("deep_real", "novel", "exceptions", "counterfactual", "ablation", "undetermined", "undet_control")
    keep = ["id", "block", "question_en", "question_es", "about", "queried predication", "scenario / premise",
            "fluency_en (1-5)", "fluency_es (1-5)"]
    wb = Workbook()
    ws = wb.active
    ws.title = "faithfulness"
    ws.append(keep + V2_COLS)
    salmon = PatternFill("solid", fgColor="F8CBAD")
    for row in src.iter_rows(min_row=2, values_only=True):
        d = dict(zip(head, row))
        if d.get("block") not in deep:
            continue
        ws.append([d.get(k) for k in keep] + [""] * len(V2_COLS))
        if "step 2:" in str(d.get("queried predication") or ""):
            for cell in ws[ws.max_row][: len(keep)]:
                cell.fill = salmon
    widths = {"id": 24, "block": 13, "question_en": 50, "question_es": 50, "about": 20, "queried predication": 40,
              "scenario / premise": 40, "comment v2": 40}
    for k, col in enumerate(keep + V2_COLS, start=1):
        cell = ws.cell(row=1, column=k)
        ws.column_dimensions[cell.column_letter].width = widths.get(col, 14)
        cell.font = Font(bold=True)
        if col in V2_COLS:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    info = wb.create_sheet("instructions", 0)
    for k, v in V2_INSTRUCTIONS:
        info.append([k, v])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out)
    print(f"{ws.max_row - 1} deep items -> {out}")


def summarise_v2(a_path: str, b_path: str):
    from openpyxl import load_workbook
    from sklearn.metrics import cohen_kappa_score

    def load(p):
        ws = load_workbook(p, read_only=True)["faithfulness"]
        rows = list(ws.iter_rows(values_only=True))
        return {r[0]: dict(zip(rows[0], r)) for r in rows[1:] if r[0]}

    a, b = load(a_path), load(b_path)
    for col in V2_COLS[:2]:
        for name, r in (("rater 1", a), ("rater 2", b)):
            v = [_num(x[col]) for x in r.values() if _num(x[col]) is not None]
            if v:
                print(f"{col} {name}: {sum(v):.0f}/{len(v)} ({100 * sum(v) / len(v):.1f} %)")
        pairs = [(_num(a[i][col]), _num(b[i][col])) for i in a if i in b]
        pairs = [(x, y) for x, y in pairs if x is not None and y is not None]
        if len(pairs) >= 5:
            x, y = zip(*pairs)
            agree = sum(p == q for p, q in pairs) / len(pairs)
            print(f"{col}: agreement {100 * agree:.1f} %, Cohen's kappa {cohen_kappa_score(x, y):.2f} (n={len(pairs)})")


if __name__ == "__main__":
    main()
